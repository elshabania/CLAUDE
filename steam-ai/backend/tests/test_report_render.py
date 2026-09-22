"""Rendering tests: DOCX opens and carries headings and DRAFT; PDF is valid and multi-page."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document
from fixtures_report import report_data_root, report_store  # noqa: F401

from steam_ai import audit
from steam_ai.reports import diagnostic
from steam_ai.reports.docx_renderer import render_docx
from steam_ai.reports.model import ReportDoc, build_diagnostic_report
from steam_ai.reports.pdf_renderer import render_pdf


@pytest.fixture
def report(report_store) -> ReportDoc:  # noqa: F811
    return build_diagnostic_report(report_store)


def _pdf_page_count(data: bytes) -> int | None:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return None
    import io

    return len(PdfReader(io.BytesIO(data)).pages)


def test_docx_renders_with_headings_and_draft(report: ReportDoc, tmp_path: Path) -> None:
    out = render_docx(report, tmp_path / "report.docx")
    assert out.exists() and out.stat().st_size > 10_000
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Executive summary" in headings
    assert "Findings by severity" in headings
    assert "Appendix" in headings
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "DRAFT" in body
    assert report.metadata.synthetic_banner in body
    header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
    assert report.metadata.watermark in header_text
    footer_xml = doc.sections[0].footer._element.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert len(doc.tables) >= 10
    assert len(doc.inline_shapes) == 4
    # source references render as small grey runs
    grey_runs = [
        r
        for p in doc.paragraphs
        for r in p.runs
        if r.text.startswith("[src: ") and r.font.size is not None and r.font.size.pt == 7
    ]
    assert len(grey_runs) > 10
    assert all(str(r.font.color.rgb) == "808080" for r in grey_runs)
    fonts = {r.font.name for p in doc.paragraphs for r in p.runs if r.font.name}
    assert fonts == {"Calibri"}


def test_pdf_renders_multiple_pages(report: ReportDoc, tmp_path: Path) -> None:
    out = render_pdf(report, tmp_path / "report.pdf")
    data = out.read_bytes()
    assert data.startswith(b"%PDF")
    assert data.rstrip().endswith(b"%%EOF")
    pages = _pdf_page_count(data)
    if pages is None:
        # reportlab writes page objects uncompressed; count them directly
        pages = len(re.findall(rb"/Type\s*/Page[^s]", data))
    assert pages > 1
    assert len(data) > 50_000


def test_diagnostic_build_writes_both_and_audits(report_store) -> None:  # noqa: F811
    out = diagnostic.build(report_store)
    assert set(out) == {"docx", "pdf"}
    assert out["docx"] == report_store.derived_path("report.docx")
    assert out["pdf"] == report_store.derived_path("report.pdf")
    assert out["docx"].exists() and out["pdf"].exists()
    events = [e for e in audit.tail(20) if e["event"] == "report_generated"]
    assert events and events[-1]["run_id"] == report_store.run_id
    assert events[-1]["n_findings"] == 12
    assert set(events[-1]["figures"]) == {
        "convergence",
        "findings_by_check",
        "network_map",
        "vc_hist",
    }


def test_diagnostic_build_custom_out_dir(report_store, tmp_path: Path) -> None:  # noqa: F811
    out = diagnostic.build(report_store, out_dir=tmp_path / "reports")
    assert out["docx"].parent == tmp_path / "reports"
    assert out["pdf"].read_bytes().startswith(b"%PDF")
