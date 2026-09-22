"""Render a :class:`ReportDoc` to DOCX with python-docx.

One typeface (Calibri), draft watermark in the header of every page, a red
DRAFT callout on the first page, shaded synthetic banner, tables with a
shaded repeating header row, figures at 16 cm, and source references as
small grey text under the paragraph or table they support.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph

from .model import (
    BulletList,
    Callout,
    Figure,
    Heading,
    PageBreak,
    Paragraph,
    ReportDoc,
    Table,
)

FONT = "Calibri"
GREY = RGBColor(0x80, 0x80, 0x80)
RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x22, 0x22, 0x22)
HEADER_FILL = "D9E2F3"
CALLOUT_FILL = {
    "draft": "FDE9E9",
    "synthetic": "FFF4CC",
    "Critical": "F8D7D7",
    "High": "FCE4D6",
    "Medium": "FFF2CC",
    "Info": "EDEDED",
    "note": "EDEDED",
}
CALLOUT_LABEL = {"draft": "DRAFT", "synthetic": "SYNTHETIC DATA"}


def _set_run_font(run, size_pt: float | None = None, colour: RGBColor | None = None,
                  bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if colour is not None:
        run.font.color.rgb = colour
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def _shade(element, fill: str) -> None:
    """Apply background shading to a paragraph (pPr) or cell (tcPr)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _style_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    normal.paragraph_format.space_after = Pt(6)
    sizes = {"Title": 26, "Subtitle": 13, "Heading 1": 16, "Heading 2": 13, "Heading 3": 11}
    for name, size in sizes.items():
        st = styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F) if name.startswith("Heading") else INK
        st.font.bold = name != "Subtitle"
        st.font.italic = False
        st_rpr = st.element.get_or_add_rPr()
        st_rfonts = st_rpr.find(qn("w:rFonts"))
        if st_rfonts is None:
            st_rfonts = OxmlElement("w:rFonts")
            st_rpr.append(st_rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            st_rfonts.set(qn(attr), FONT)
        # remove theme font references so Calibri wins
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if st_rfonts.get(qn(attr)) is not None:
                del st_rfonts.attrib[qn(attr)]


def _add_field(paragraph: DocxParagraph, instr: str) -> None:
    run = paragraph.add_run()
    _set_run_font(run, 8, GREY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = f" {instr} "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, text, sep, placeholder, end):
        run._element.append(el)


def _header_footer(doc: Document, report: ReportDoc) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.different_first_page_header_footer = False

    header = section.header
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run(report.metadata.watermark)
    _set_run_font(run, 8, GREY, italic=True)

    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(
        f"{report.title} | run {report.metadata.run_id} | generated "
        f"{report.metadata.generated_at} by STEAM-AI {report.metadata.steam_ai_version} | Page "
    )
    _set_run_font(r, 8, GREY)
    _add_field(fp, "PAGE")
    r2 = fp.add_run(" of ")
    _set_run_font(r2, 8, GREY)
    _add_field(fp, "NUMPAGES")


def _source_ref(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_run_font(run, 7, GREY)


def _callout(doc: Document, block: Callout) -> None:
    p = doc.add_paragraph()
    _shade(p._p.get_or_add_pPr(), CALLOUT_FILL.get(block.severity, CALLOUT_FILL["note"]))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    label = CALLOUT_LABEL.get(block.severity, block.severity)
    lr = p.add_run(f"{label}: ")
    _set_run_font(lr, 10.5, RED if block.severity in ("draft", "Critical") else INK, bold=True)
    tr = p.add_run(block.text)
    _set_run_font(tr, 10.5, INK)


def _table(doc: Document, block: Table) -> None:
    if block.caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(2)
        cr = cp.add_run(block.caption)
        _set_run_font(cr, 9.5, INK, bold=True)
    n_cols = len(block.columns)
    tbl = doc.add_table(rows=1, cols=n_cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for cell, col in zip(hdr.cells, block.columns):
        _shade(cell._tc.get_or_add_tcPr(), HEADER_FILL)
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(col))
        _set_run_font(run, 9, INK, bold=True)
    for row in block.rows:
        cells = tbl.add_row().cells
        for cell, value in zip(cells, list(row) + [""] * (n_cols - len(row))):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(value))
            _set_run_font(run, 9, INK)
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
    if block.source_ref:
        _source_ref(doc, block.source_ref)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _figure(doc: Document, block: Figure) -> None:
    path = Path(block.path_png)
    if not path.exists():
        p = doc.add_paragraph(f"Figure not available: {path.name}")
        _set_run_font(p.runs[0], 9, GREY, italic=True)
        return
    doc.add_picture(str(path), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(2)
    run = cp.add_run(block.caption)
    _set_run_font(run, 9, INK, italic=True)
    if block.source_ref:
        _source_ref(doc, block.source_ref)


def _cover(doc: Document, report: ReportDoc) -> None:
    t = doc.add_paragraph(style="Title")
    tr = t.add_run(report.title)
    _set_run_font(tr, 26, INK, bold=True)
    s = doc.add_paragraph(style="Subtitle")
    sr = s.add_run(report.subtitle)
    _set_run_font(sr, 13, INK)
    meta = report.metadata
    m = doc.add_paragraph()
    mr = m.add_run(
        f"Generated {meta.generated_at} by STEAM-AI {meta.steam_ai_version}. "
        f"Ingested {meta.ingested_at}."
    )
    _set_run_font(mr, 9, GREY)


def render_docx(report: ReportDoc, out_path: Path) -> Path:
    doc = Document()
    _style_document(doc)
    _header_footer(doc, report)
    _cover(doc, report)

    for block in report.blocks:
        if isinstance(block, Heading):
            level = min(max(block.level, 1), 3)
            h = doc.add_heading(level=level)
            run = h.add_run(block.text)
            _set_run_font(run, {1: 16, 2: 13, 3: 11}[level], RGBColor(0x1F, 0x3A, 0x5F),
                          bold=True)
        elif isinstance(block, Paragraph):
            p = doc.add_paragraph()
            if block.source_ref:
                p.paragraph_format.space_after = Pt(0)
            run = p.add_run(block.text)
            _set_run_font(run, 10.5, INK)
            if block.source_ref:
                _source_ref(doc, block.source_ref)
        elif isinstance(block, BulletList):
            for item in block.items:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(item)
                _set_run_font(run, 10.5, INK)
        elif isinstance(block, Table):
            _table(doc, block)
        elif isinstance(block, Figure):
            _figure(doc, block)
        elif isinstance(block, PageBreak):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(block, Callout):
            _callout(doc, block)
        else:  # pragma: no cover - defensive
            raise TypeError(f"unknown block type {type(block).__name__}")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
